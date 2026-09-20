#!/usr/bin/env python3
"""Completa la plantilla PPTX manteniendo su slide, tema y dimensiones."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


def replace_value(paragraph, value: str) -> None:
    """Conserva el primer run (etiqueta) y reemplaza el resto por el valor."""
    if not paragraph.runs:
        paragraph.add_run().text = value
        return
    label = paragraph.runs[0].text.rstrip().rstrip(":")
    paragraph.runs[0].text = f"{label}:"
    if len(paragraph.runs) == 1:
        paragraph.add_run().text = f" {value}"
    else:
        paragraph.runs[1].text = f" {value}"
        for run in paragraph.runs[2:]:
            run.text = ""


def update_text(slide, metrics: dict) -> None:
    replacements = {
        "Hora de inicio": metrics["start_display"],
        "Hora de finalización": metrics["end_display"],
        "Duración": metrics["duration_display"],
        "Cantidad total de transacciones": str(metrics["transactions_total"]),
        "Transacciones exitosas y fallidas": (
            f'{metrics["transactions_successful"]} exitosas y '
            f'{metrics["transactions_failed"]} fallidas'
        ),
        "TPS promedio": f'{metrics["tps_average"]:.2f} TPS',
        "TPS máximo": f'{metrics["tps_maximum"]} TPS',
    }

    found = set()
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            normalized = paragraph.text.strip()
            for label, value in replacements.items():
                if normalized.startswith(label):
                    replace_value(paragraph, value)
                    found.add(label)
                    break
            if normalized.startswith("Gráfica de usuarios activos durante el ramp-up"):
                paragraph.text = "Gráfica de usuarios activos y TPS durante la prueba:"

    missing = set(replacements).difference(found)
    if missing:
        raise ValueError(f"No se encontraron campos en la plantilla: {sorted(missing)}")


def replace_chart(slide, chart_path: Path) -> None:
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if not pictures:
        raise ValueError("La plantilla no contiene una imagen que pueda reemplazarse por la gráfica")
    picture = pictures[0]
    left, top, width, height = picture.left, picture.top, picture.width, picture.height
    # La plantilla recibida reserva una caja pequeña. Se amplía dentro del espacio
    # libre del mismo slide para que ambas escalas resulten legibles.
    if width < Inches(7):
        left = Inches(1.75)
        top = Inches(5.05)
        width = Inches(9.75)
        height = Inches(2.05)
    picture._element.getparent().remove(picture._element)
    slide.shapes.add_picture(str(chart_path), left, top, width=width, height=height)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--metrics", required=True, type=Path)
    parser.add_argument("--chart", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    metrics = json.loads(args.metrics.read_text(encoding="utf-8"))
    presentation = Presentation(args.template)
    if not presentation.slides:
        raise ValueError("La plantilla no contiene slides")

    slide = presentation.slides[0]
    update_text(slide, metrics)
    replace_chart(slide, args.chart)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(args.output)
    print(f"Reporte generado: {args.output}")


if __name__ == "__main__":
    main()
