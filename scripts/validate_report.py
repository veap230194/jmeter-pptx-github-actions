#!/usr/bin/env python3
"""Validación estructural mínima del reporte generado."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("report", type=Path)
    args = parser.parse_args()

    presentation = Presentation(args.report)
    assert len(presentation.slides) == 1, "El reporte debe conservar un slide"
    slide = presentation.slides[0]
    all_text = "\n".join(
        shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    )
    required = [
        "Hora de inicio", "Hora de finalización", "Duración",
        "Cantidad total de transacciones", "TPS promedio", "TPS máximo"
    ]
    missing = [text for text in required if text not in all_text]
    assert not missing, f"Faltan textos obligatorios: {missing}"
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes), (
        "El reporte no contiene la gráfica"
    )
    print("Validación estructural correcta")


if __name__ == "__main__":
    main()

